"""Fill the teacher's visual template with the FerreMax architecture report.

The original guide stays unchanged. Its instructions and examples are not
copied into the students' report. Render with the documents skill after build.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import re,json,html
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import fitz

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT.parent/'outputs';OUT.mkdir(exist_ok=True)
TMP=ROOT.parent/'tmp'/'guide_report';TMP.mkdir(exist_ok=True,parents=True)
SOURCE=ROOT/'reference'/'Actividad_Unidad_3_CLN_ArquitecturaAlumnos(1).docx'
DATA=json.loads((ROOT/'docs'/'current_sources.json').read_text())
doc=Document(SOURCE)
front=[deepcopy(p._p) for p in doc.paragraphs[:4]]
meta=deepcopy(doc.tables[0]._tbl)
for e in list(doc.element.body):
    if e.tag!=qn('w:sectPr'):doc.element.body.remove(e)

# Keep the guide's Letter layout, margins, theme, and restrained typography.
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    doc.styles[name].paragraph_format.keep_with_next=True
    pp=doc.styles[name].element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
normal=doc.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10.5)
for section in doc.sections:
    for story in [section.header,section.footer,section.first_page_header,section.first_page_footer,section.even_page_header,section.even_page_footer]:
        for t in story._element.xpath('.//w:t'):
            if t.text:
                t.text=t.text.replace('Cisco Confidential','').replace('Cisco Confidencial','').replace('Pagina ','Página ').replace('Componentes en Logica de Negocios','Componentes de lógica de negocio')
doc.core_properties.title='Diseño arquitectural del sistema con IA FerreMax'
doc.core_properties.subject='Componentes de lógica de negocio Unidad 3'
doc.core_properties.author='Equipo 17'
doc.core_properties.last_modified_by='Equipo 17'
doc.core_properties.comments=''

def clean(t):
    t=re.sub(r'<mention-date start="([^"]+)"[^>]*/>',r'\1',t)
    return html.unescape(re.sub(r'<[^>]+>','',t)).replace('–','-').replace('—','-').replace('\u2011','-')

def runs(p,t):
    t=clean(t)
    t=re.sub(r'\[([^\]]+\.(?:py|md|sql))\]\(https?://[^)]+\)',r'`\1`',t)
    for piece in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\(https?[^)]+\))',t):
        if not piece:continue
        if piece.startswith('**'):
            r=p.add_run(piece[2:-2]);r.bold=True
        elif piece.startswith('`'):
            r=p.add_run(piece[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        elif re.match(r'\[.*\]\(https?',piece):
            m=re.match(r'\[([^\]]+)\]\(([^)]+)\)',piece)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'0563C1');pr.append(co)
            u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);text=OxmlElement('w:t');text.text=m[1];r.append(text);link.append(r);p._p.append(link)
        else:p.add_run(piece)
    return p

def para(t,style=None):
    p=doc.add_paragraph(style=style);runs(p,t);return p

def heading(t,level=1):
    return doc.add_paragraph(clean(t),style='Heading '+str(level))

def rows_from(xml):
    return [[re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',clean(c)).replace('**','').replace('`','') for c in re.findall(r'<td[^>]*>(.*?)</td>',r,re.S)] for r in re.findall(r'<tr[^>]*>(.*?)</tr>',xml,re.S)]

def table(rows,widths=None,blue=False):
    if not rows:return
    n=len(rows[0]);weights={2:[.29,.71],3:[.23,.35,.42],4:[.19,.25,.25,.31],5:[.16,.23,.23,.24,.14],6:[.10,.16,.16,.23,.15,.20],7:[.07,.14,.10,.12,.21,.10,.26]}
    if widths is None:
        widths=weights.get(n,[1/n]*n)
        if rows[0][0]=='Regla':widths=[.075,.16,.16,.24,.365]
        elif rows[0][0]=='Principio':widths=[.13,.31,.25,.31]
        elif rows[0][0]=='Tipo':widths=[.13,.20,.20,.23,.24]
        elif rows[0][0]=='Necesidad':rows[0][0]='ID';rows[0][5]='ADR';widths=[.08,.15,.12,.13,.20,.11,.21]
        elif rows[0][0]=='ID':widths=[.07,.18,.15,.25,.13,.22]
        elif rows[0][0]=='Servicio':widths=[.15,.14,.20,.23,.12,.16]
        elif rows[0][0]=='Iteración':widths=[.12,.13,.21,.19,.18,.17]
    widths=[6.85*x/sum(widths) for x in widths]
    t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.LEFT;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    props.append(borders)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trPr=cells[0]._tc.getparent().get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0:trPr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),('2E75B6' if blue else '0D2238') if i==0 else ('ECF4FA' if i%2==0 else 'FFFFFF'));cp.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                mm=OxmlElement('w:'+edge);mm.set(qn('w:w'),'85' if edge in ['top','bottom'] else '95');mm.set(qn('w:type'),'dxa');margins.append(mm)
            cp.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=False
            if i>0 and re.fullmatch(r'(CU\d+|BN\d+|BR\d+|ADR\d+|[SOLID]|\d+)',text):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(text);r.font.name='Aptos';r.font.size=Pt(9.5);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
    return t

def md(t,blue=False,sublevel=2):
    for c in re.split(r'(<table\b.*?</table>|```.*?```)',t,flags=re.S):
        if c.startswith('<table'):table(rows_from(c),blue=blue);continue
        if c.startswith('```'):
            for s in c.splitlines()[1:-1]:
                p=para(s);p.paragraph_format.space_after=Pt(0)
                for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
            continue
        for s in c.splitlines():
            s=s.strip()
            if not s:continue
            if s.startswith('### '):heading(s[4:],sublevel)
            elif s.startswith('## '):heading(s[3:],1)
            elif s.startswith('- '):para(s[2:],'List Bullet')
            else:para(s)

def body(section):return section.split('\n',1)[1].strip()
sections={int(re.match(r'## (\d+)',s)[1]):s for s in re.split(r'(?=^## )',DATA['report'],flags=re.M) if re.match(r'## \d+',s)}
decisions={s.splitlines()[0][3:]:s for s in re.split(r'(?=^## )',DATA['decisions'],flags=re.M) if s.startswith('## ')}

def cover():
    for e in front:doc.element.body.insert(len(doc.element.body)-1,e)
    for p,text in zip(doc.paragraphs[:4],[
        'ACTIVIDAD PRÁCTICA',
        'Diseño arquitectural del sistema con IA',
        'FerreMax Informe de arquitectura',
        'FerreMax requiere coordinar cotizaciones, pedidos e inventario sin comprometer mercancía inexistente. La arquitectura propuesta organiza las reglas del negocio con Domain Model, una Service Layer y adaptadores de persistencia. El recorrido detallado cubre la consulta de disponibilidad, la revalidación de cotizaciones, la conversión a pedido, la confirmación con reserva y la cancelación antes del surtido.'
    ]):
        # Retain paragraph and first-run formatting from the actual guide.
        if p.runs:
            p.runs[0].text=text
            for r in p.runs[1:]:r.text=''
        else:p.add_run(text)
    doc.element.body.insert(len(doc.element.body)-1,meta)
    from docx.table import Table
    t=Table(meta,doc._body)
    vals=[('Empresa asignada','FerreMax'),('Integrantes','Roberto Santana Trujillo  2227084\nIsaías Leonardo Salazar Torres  2227016'),('Fecha de entrega','Revisión 16 de octubre de 2026\nEntrega 17 de octubre de 2026 a las 23:59')]
    for row,(label,value) in zip(t.rows,vals):
        for cell,text in zip(row.cells,[label,value]):
            p=cell.paragraphs[0]
            if p.runs:p.runs[0].text=text
            else:p.add_run(text)
            for r in p.runs[1:]:r.text=''
            for r in p.runs:r.font.size=Pt(10)
    para('Profesor Miguel Ángel Gómez Marroquín')
    doc.add_page_break()

def figure(id,title,first=False):
    if not first:doc.add_page_break()
    heading(title,2)
    p=fitz.open(ROOT/'diagrams'/(id+'.pdf'))[0]
    bounds=json.loads((ROOT/'diagrams'/(id+'_segments.json')).read_text()) if id in ['D04','D05'] else [0,p.rect.height]
    for n,(top,bottom) in enumerate(zip(bounds,bounds[1:])):
        if n:
            doc.add_page_break();heading(title+' continuación',2)
            para('Participantes de izquierda a derecha: Caja, PedidoService, PostgresUnitOfWork, PgPedidos y PgInventarios.')
        rect=fitz.Rect(0,top,p.rect.width,bottom);path=TMP/(id+'_'+str(n)+'.png');p.get_pixmap(matrix=fitz.Matrix(1.8,1.8),clip=rect,alpha=False).save(path)
        width=min(6.85,8.15*rect.width/rect.height);pic=doc.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(path),width=Inches(width))
        cap=para('Fuente '+id+'.mmd y exportación '+id+'.pdf en diagrams.');cap.paragraph_format.space_after=Pt(6)
        for r in cap.runs:r.font.size=Pt(9)
        if id=='D04':
            para('El servicio valida el rol, obtiene el pedido y bloquea el inventario por SKU en un orden uniforme. Pedido e InventarioUbicacion aplican sus reglas; la confirmación, las reservas, la auditoría y el evento se guardan juntos. La respuesta Confirmado se emite después del COMMIT.' if n==0 else 'La outbox conserva PedidoConfirmado dentro de la misma transacción. Su registro acredita el hecho del negocio; el consumidor externo que lo publicará es una ampliación posterior.')
        if id=='D05':
            para('La consulta previa de disponibilidad es informativa. Al confirmar se recupera el inventario actual con bloqueo; si otro pedido consumió la última unidad, la operación detecta disponibilidad insuficiente.' if n==0 else 'El rollback descarta la confirmación y todas las reservas provisionales. El pedido permanece Pendiente. El tratamiento financiero de un pago autorizado sin reserva sigue siendo la política pendiente H05; no se presupone un reembolso automático.')

def main():
    cover()
    heading('1 Preparación del caso')
    intro=body(sections[1]).replace('La guía permite código parcial; por ello, cada sección distingue diseño de negocio, implementación existente y trabajo posterior.','El prototipo representa este subconjunto; los procesos posteriores se conservan en el mapa general del dominio.')
    md(intro)
    heading('Continuidad con la Unidad 2',2)
    sources=body(sections[2]).replace('BPMN cotejado visualmente en el entorno cloud.','BPMN cotejado visualmente.')
    # Course instructions are references, not rows of the business evidence table.
    first=re.search(r'<table.*?</table>',sources,re.S)[0];rr=rows_from(first)
    table(rr[:5])
    md(sources.split('</table>',1)[1])
    heading('Alcance y casos de uso',2);md(body(sections[3]),blue=True)
    doc.add_page_break();heading('2 Selección del patrón de lógica de negocio')
    comparison=body(sections[4]);idx=comparison.index('### Evaluación de Transaction Script')
    md(comparison[:idx],blue=True);doc.add_page_break()
    narrative=comparison[idx:]
    narrative=re.sub(r'\*\*Estado de la decisión:\*\*[^\n]*','',narrative)
    md(narrative)
    doc.add_page_break();heading('3 Mapa de dominio')
    domain=body(sections[5]);xml=re.search(r'<table.*?</table>',domain,re.S)[0];rr=rows_from(xml)
    para('El mapa general identifica siete agregados. Cada raíz controla sus invariantes y conserva referencias por identidad hacia otros agregados.')
    for row in rr[1:]:
        heading(row[0],3)
        for label,value in zip(rr[0][1:],row[1:]):para('**'+label+'** '+value)
    md(domain.split('</table>',1)[1])
    heading('4 Service Layer');md(body(sections[6]),blue=True)
    heading('5 Persistencia y transacciones');md(body(sections[7]))
    heading('PostgreSQL en Supabase',2)
    db=body(sections[10]);db=re.sub(r'\*\*Estado:\*\*[^\n]*','El adaptador PostgreSQL y el esquema están preparados; su prueba con conexiones reales sigue pendiente.',db)
    md(db)
    heading('6 Commands y Queries');md(body(sections[8]),blue=True)
    heading('7 Revisión con SOLID');md(body(sections[9]),blue=True)
    doc.add_page_break();heading('8 Representación de la arquitectura')
    para('D01 presenta el mapa general. D02 describe las clases del dominio y D03 las de aplicación e infraestructura. D04 recorre la confirmación con reserva y D05 muestra el rechazo por falta de disponibilidad. Los nombres corresponden con los componentes del código.')
    figure('D01','D01 Mapa general del dominio',True)
    for id,title in [('D02','D02 Clases del dominio'),('D03','D03 Aplicación e infraestructura'),('D04','D04 Confirmación y reserva'),('D05','D05 Rechazo por disponibilidad insuficiente')]:figure(id,title)
    doc.add_page_break();heading('Registros de decisiones arquitecturales',2)
    for name,section in decisions.items():
        if name.startswith('ADR'):
            heading(name,2);md(body(section),sublevel=3)
    heading('9 Validación y trazabilidad')
    validation=body(sections[11]);xml=re.search(r'<table.*?</table>',validation,re.S)[0];rr=rows_from(xml)
    rr=[r for r in rr if r[0] not in ['Paquete y GitHub','GitHub y PDF FIME']]
    md(validation.split('<table',1)[0]);table(rr)
    for name in ['Matriz de trazabilidad por necesidad','Matriz de trazabilidad por regla','Trazabilidad de los rechazos']:
        heading(name,2);md(body(decisions[name]))
    heading('Evidencia de ejecución local',2)
    para('El ensayo ejecutó 21 pruebas: 20 correctas y una integración PostgreSQL omitida por falta de TEST_DATABASE_URL. Los registros completos están en docs/evidence. Memory protege la competencia con un bloqueo de proceso; el resultado no certifica los bloqueos ni el rendimiento de PostgreSQL.')
    for line in (ROOT/'docs/evidence/pruebas_locales.txt').read_text().splitlines()[-4:]:
        p=para(line);p.paragraph_format.space_after=Pt(0)
        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
    heading('10 Uso de IA')
    para('La IA apoyó la comparación de alternativas, el código y la documentación. La bitácora distingue las decisiones expresamente confirmadas por Roberto de las propuestas técnicas. La revisión conjunta del equipo y la defensa no se registran como realizadas.')
    md(body(decisions['Bitácora de IA']))
    heading('Correcciones posteriores',2)
    md(body(decisions['Iteraciones posteriores verificadas']))
    para('**IA10** Se generaron y revisaron las exportaciones y el paquete del proyecto. **IA11** Se adaptó el informe al formato de la guía práctica, conservando el contenido de FerreMax y retirando instrucciones, rúbrica y ejemplos ajenos al proyecto.')
    heading('11 Conclusiones')
    conclusions=body(sections[12]);conclusions='\n'.join(p for p in conclusions.splitlines() if not p.startswith('Estas conclusiones son'))
    md(conclusions)
    first_ref=len(doc.paragraphs)
    heading('12 Referencias');md(body(sections[13]))
    for p in doc.paragraphs[first_ref:-1]:p.paragraph_format.keep_with_next=True
    path=OUT/'FerreMax_U3_Informe.docx';doc.save(path)
    (ROOT/'docs'/'FerreMax_U3_Informe.docx').write_bytes(path.read_bytes())
    # Keep the readable repository copy aligned with the actual Word report.
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    lines=[];figure_ids=iter(['D01','D02','D03','D04','D04','D05','D05']);last_figure=None
    for element in doc.element.body:
        if element.tag==qn('w:p'):
            p=Paragraph(element,doc._body)
            if p._p.xpath('.//w:drawing'):
                id=next(figure_ids)
                if id!=last_figure:lines.append('!['+id+'](../diagrams/'+id+'.png)\n')
                last_figure=id
            words=[]
            for child in element:
                text=''.join(child.xpath('.//w:t/text()'))
                if not text:continue
                if child.tag==qn('w:hyperlink'):
                    target=p.part.rels[child.get(qn('r:id'))].target_ref;words.append('['+text+']('+target+')')
                else:words.append(text)
            text=''.join(words).strip()
            if not text:continue
            style=p.style.name
            if style=='Title':text='# '+text
            elif style.startswith('Heading '):text='#'*(int(style[-1])+1)+' '+text
            lines.append(text+'\n')
        elif element.tag==qn('w:tbl'):
            t=Table(element,doc._body)
            for i,row in enumerate(t.rows):
                values=[cell.text.replace('|','\\|').replace('\n','<br>') for cell in row.cells]
                lines.append('| '+' | '.join(values)+' |')
                if i==0:lines.append('| '+' | '.join(['---']*len(values))+' |')
            lines.append('')
    (ROOT/'docs'/'Informe_Arquitectura.md').write_text('\n'.join(lines))
    print(path)
if __name__=='__main__':main()
